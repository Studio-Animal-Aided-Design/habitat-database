#!/usr/bin/env python3
"""Build the editable operator handbook from the three canonical Markdown chapters."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCES = ("installation.md", "konfiguration.md", "betrieb.md")
OUTPUT = ROOT / "handbuch.docx"


def plain(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("`", "")
    return text.replace("**", "")


def shade(element, fill: str) -> None:
    properties = element.get_or_add_tcPr() if element.tag == qn("w:tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    properties.append(shd)


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)


def set_cell_margin(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def apply_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(25, 32, 40)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title_properties = title._element.pPr
    if title_properties is not None:
        border = title_properties.find(qn("w:pBdr"))
        if border is not None:
            title_properties.remove(border)

    for name, size, before, after in (("Heading 1", 16, 18, 8), ("Heading 2", 12, 13, 5), ("Heading 3", 10.5, 10, 3)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = doc.styles.add_style("Codeblock", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor(25, 32, 40)
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_after = Pt(2)
    code.paragraph_format.line_spacing = 1.05

    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.size = Pt(10.5)
        doc.styles[name].paragraph_format.space_after = Pt(3)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows[0]) == 5:
        for row in rows[1:]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.keep_with_next = True
            p.add_run(plain(row[0])).bold = True
            meta = doc.add_paragraph()
            meta.paragraph_format.left_indent = Inches(0.15)
            meta.paragraph_format.space_after = Pt(2)
            meta.paragraph_format.keep_with_next = True
            run = meta.add_run(f"{plain(row[1])}  ·  #105: {plain(row[2])}  ·  #117: {plain(row[3])}")
            run.font.size = Pt(9)
            detail = doc.add_paragraph(plain(row[4]))
            detail.paragraph_format.left_indent = Inches(0.15)
            detail.paragraph_format.space_after = Pt(4)
            detail.paragraph_format.keep_together = True
        return

    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.autofit = False
    for row_number, values in enumerate(rows):
        cells = table.rows[0].cells if row_number == 0 else table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = plain(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margin(cell)
            if row_number == 0:
                shade(cell._tc, "E7EEF3")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_chapter(doc: Document, path: Path, first: bool) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    paragraph_lines: list[str] = []
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    in_code = False

    def flush_paragraph() -> None:
        if paragraph_lines:
            doc.add_paragraph(plain(" ".join(paragraph_lines)))
            paragraph_lines.clear()

    def flush_table() -> None:
        if table_rows:
            add_table(doc, table_rows)
            table_rows.clear()

    for line in lines + [""]:
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            flush_table()
            if in_code:
                for code_index, code_line in enumerate(code_lines):
                    p = doc.add_paragraph(style="Codeblock")
                    p.add_run(code_line)
                    shade(p._p, "F3F5F6")
                    p.paragraph_format.keep_with_next = code_index < len(code_lines) - 1
                    if code_index == len(code_lines) - 1:
                        p.paragraph_format.space_after = Pt(8)
                code_lines.clear()
                in_code = False
            else:
                if doc.paragraphs:
                    doc.paragraphs[-1].paragraph_format.keep_with_next = True
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            cells = [part.strip() for part in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            table_rows.append(cells)
            continue
        flush_table()
        if not stripped:
            flush_paragraph()
            continue
        if stripped.startswith("#"):
            flush_paragraph()
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_paragraph(plain(stripped[level:].strip()), style=f"Heading {min(level, 3)}")
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            doc.add_paragraph(plain(stripped[2:]), style="List Bullet")
            continue
        number = re.match(r"^\d+\.\s+", stripped)
        if number:
            flush_paragraph()
            doc.add_paragraph(plain(stripped[number.end():]), style="List Number")
            continue
        paragraph_lines.append(stripped)


def main() -> None:
    doc = Document()
    apply_styles(doc)
    doc.core_properties.title = "Habitat Datenbank Installation Konfiguration und Betrieb"
    doc.core_properties.subject = "Handbuch für IONOS Bereitstellung und Betrieb"
    doc.core_properties.keywords = "Habitat Datenbank, IONOS, Deployment, Betrieb"

    doc.add_paragraph("Habitat Datenbank Installation Konfiguration und Betrieb", style="Title")
    intro = doc.add_paragraph(
        "Schrittweise Anleitung für die erstmalige Bereitstellung und den späteren Betrieb "
        "der Replacement App auf einem IONOS Server. Die drei Kapitel unterscheiden "
        "Installationsschritte, externe Entscheidungen und Laufzeitkonfiguration sowie "
        "den laufenden Betrieb. Stand 20. September 2026."
    )
    intro.paragraph_format.space_after = Pt(10)

    for index, source in enumerate(SOURCES):
        add_chapter(doc, ROOT / source, first=index == 0)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
